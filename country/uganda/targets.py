"""The Ministry of Health's October 2026 ICHD targets as ICRTargetPopulation Groups.

    python3 -I targets.py "Oct 2026_ICHDs_EPI Targets.xlsx" build/snapshot/locations.ndjson build/targets.ndjson

Reads the workbook's Target_refined sheet (the values Excel saved, not the formulas) and writes,
per target column, one Group for each of the 146 districts and cities, plus region and country
totals (sums of the district figures, is-calculated). Also writes the two code systems the Groups
use: the target columns (Group.code) and the round (meta.tag), which holds the round's dates
until the campaign itself is stored.

The Groups are the targets a campaign will plan against (is-planning-denominator), except the
two population columns that the targets are derived from. Counts are rounded to whole people.
"""
import html, json, sys

import openpyxl

XLSX, LOCATIONS, OUT = sys.argv[1:4]
SERVER = "https://fhir-uga.healthcampaigns.org/fhir"
TARGET_CS, ROUND_CS = f"{SERVER}/CodeSystem/uga-ichd-target", f"{SERVER}/CodeSystem/uga-campaign-round"
ICR = "https://icr.healthcampaigns.org"
SD, GROUP_CHAR = f"{ICR}/StructureDefinition", f"{ICR}/CodeSystem/icr-group-characteristic-cs"
SOURCE_CS = f"{ICR}/CodeSystem/icr-denominator-source-cs"
ROUND, ROUND_NAME, ROUND_START, ROUND_END = "ichd-2026-10", "ICHDs October 2026", "2026-10-01", "2026-10-31"
ESTIMATE_DATE = "2026-09-28"   # the workbook was last saved on this date
WORKBOOK = "Oct 2026_ICHDs_EPI Targets.xlsx, sheet Target_refined"
UNDER_1 = (0, "mo", 11, "mo")

# Workbook column → (code, display, how the workbook gets the figure, calculated, planning,
# denominator type, age band). Column letters are the sheet's; "Total target" (M) adds up
# targets of different interventions and is left out.
COLUMNS = {
    "B": ("population-2026", "Total population 2026", "Population_2026, as given.",
          False, False, "total-population", None),
    "C": ("catch-population-2025", "Catch population 2025", "Catch Population_2025, as given.",
          False, False, "total-population", None),
    "D": ("under-1", "Children under 1 year", "Catch population 2025 × 4.3%.",
          True, True, "at-risk", UNDER_1),
    "E": ("under-1-unimmunized", "Unimmunized children under 1 year", "Children under 1 year × 2%.",
          True, True, "at-risk", UNDER_1),
    "F": ("under-1-under-immunized", "Under-immunized children under 1 year",
          "Children under 1 year × 3%.", True, True, "at-risk", UNDER_1),
    "G": ("born-2025-04", "Children born in April 2025 (MR2 target, to confirm)",
          "One month's births: total population 2026 × 4.3% ÷ 12. The column is headed April 2025; "
          "these children are 18 months old in October 2026.", True, True, "at-risk", None),
    "H": ("under-1-total", "Total under-1 target",
          "Unimmunized + under-immunized + children born in April 2025.", True, True, "at-risk", None),
    "I": ("mr2-born-2024-04", "MR2 target: children born in April 2024",
          "One month's births: catch population 2025 × 4.3% ÷ 12 (column MR2 Target_April 2024).",
          True, True, "at-risk", None),
    "J": ("deworming-vitamin-a", "Deworming and vitamin A", "Deworming+VitA, as given.",
          False, True, "at-risk", None),
    "K": ("hpv", "HPV vaccine", "HPV, as given.", False, True, "at-risk", None),
    "L": ("td2-plus", "Td2+", "Td2+, as given.", False, True, "at-risk", None),
}

# Locations by the workbook's name for them (the alias fhir-pipes kept, or the name itself).
locs, by_name = {}, {}
with open(LOCATIONS) as f:
    for line in f:
        r = json.loads(line)
        locs[r["id"]] = r
        for n in [r["name"], *(r.get("alias") or [])]:
            by_name[n.strip().lower()] = r["id"]

ws = openpyxl.load_workbook(XLSX, data_only=True, read_only=True)["Target_refined"]
header = [c.value for c in next(ws.iter_rows(min_row=2, max_row=2))]
assert header[0] == "District" and header[8] == "MR2 Target_April 2024" and header[12] == "Total target", header
values = {}                      # location id → {column letter: figure}
for row in ws.iter_rows(min_row=3, values_only=True):
    if not row[0] or row[0].strip() == "National":
        continue
    lid = by_name.get(row[0].strip().lower())
    assert lid, f"no Location for {row[0]!r}"
    values[lid] = {col: row[ord(col) - ord("A")] for col in COLUMNS}
districts = [l for l, r in locs.items() if r.get("partOf") and locs[r["partOf"]["reference"][9:]].get("partOf")]
assert len(values) == 146 and set(values) == set(districts), (len(values), set(districts) ^ set(values))

# Region and country totals from the unrounded district figures.
for lid in list(values):
    for parent in (locs[lid]["partOf"]["reference"].split("/")[1], "uga"):
        total = values.setdefault(parent, dict.fromkeys(COLUMNS, 0))
        for col in COLUMNS:
            total[col] += values[lid][col]


def qty(value, unit):
    return {"value": value, "unit": {"mo": "months", "a": "years"}[unit],
            "system": "http://unitsofmeasure.org", "code": unit}


def group(lid, col):
    code, display, method, calculated, planning, dtype, age = COLUMNS[col]
    loc = locs[lid]
    summed = lid not in districts
    how = (f"The sum of its districts' figures. {method}" if summed else method)
    name = f"{display}, {loc['name']}, {ROUND_NAME}"
    chars = [{"code": {"coding": [{"system": GROUP_CHAR, "code": "geography", "display": "Geographic scope"}]},
              "valueReference": {"reference": f"Location/{lid}", "display": loc["name"]}, "exclude": False}]
    if age:
        chars.append({"code": {"coding": [{"system": GROUP_CHAR, "code": "age-band", "display": "Age band"}]},
                      "valueRange": {"low": qty(age[0], age[1]), "high": qty(age[2], age[3])}, "exclude": False})
    return {
        "resourceType": "Group", "id": f"{ROUND}-{code}-{lid}",
        "meta": {"profile": [f"{SD}/ICRTargetPopulation"],
                 "tag": [{"system": ROUND_CS, "code": ROUND, "display": ROUND_NAME}]},
        "text": {"status": "generated", "div": (
            '<div xmlns="http://www.w3.org/1999/xhtml"><p>' + html.escape(name) + "</p><p>"
            + html.escape(f"Ministry of Health targets for the {ROUND_NAME} round "
                          f"({ROUND_START} to {ROUND_END}): {WORKBOOK}. {how}") + "</p></div>")},
        "type": "person", "actual": False,
        "code": {"coding": [{"system": TARGET_CS, "code": code, "display": display}], "text": display},
        "name": name,
        "quantity": round(values[lid][col]),
        "characteristic": chars,
        "extension": [
            {"url": f"{SD}/denominator-source", "valueCodeableConcept": {
                "coding": [{"system": SOURCE_CS, "code": "govt-estimate", "display": "Government estimate"}],
                "text": f"Uganda Ministry of Health, {WORKBOOK}"}},
            {"url": f"{SD}/denominator-type", "valueCode": dtype},
            {"url": f"{SD}/estimate-date", "valueDate": ESTIMATE_DATE},
            {"url": f"{SD}/is-calculated", "valueBoolean": calculated or summed},
            {"url": f"{SD}/is-planning-denominator", "valueBoolean": planning},
        ],
    }


def codesystem(id, url, title, description, concepts, properties=()):
    rows = "".join(f"<tr><td>{c['code']}</td><td>{html.escape(c['display'])}</td></tr>" for c in concepts)
    return {"resourceType": "CodeSystem", "id": id, "url": url, "version": "2026-10",
            "name": "".join(w.capitalize() for w in id.split("-")),
            "title": title, "status": "active", "experimental": False, "publisher": "ICR", "description": description,
            "text": {"status": "generated", "div": f'<div xmlns="http://www.w3.org/1999/xhtml"><p>{html.escape(title)}</p>'
                                                   f"<table><tbody>{rows}</tbody></table></div>"},
            "caseSensitive": True, "content": "complete", "concept": concepts,
            **({"property": list(properties)} if properties else {})}


target_cs = codesystem(
    "uga-ichd-target", TARGET_CS, "Uganda ICHD target",
    "The target columns of the Ministry of Health's ICHD targets workbook. Group.code on Uganda's "
    "ICRTargetPopulation Groups, searchable with Group?code= and this system.",
    [{"code": c, "display": d, "definition": m} for c, d, m, *_ in COLUMNS.values()])
round_cs = codesystem(
    "uga-campaign-round", ROUND_CS, "Uganda campaign round",
    "Campaign rounds in Uganda, used as meta.tag on the records prepared for a round "
    "(searchable with ?_tag= and this system) before the campaign itself is stored.",
    [{"code": ROUND, "display": ROUND_NAME,
      "definition": "Integrated Child Health Days, October 2026: the whole month for now.",
      "property": [{"code": "start", "valueDateTime": ROUND_START}, {"code": "end", "valueDateTime": ROUND_END}]}],
    [{"code": "start", "description": "First day of the round", "type": "dateTime"},
     {"code": "end", "description": "Last day of the round", "type": "dateTime"}])

n = 0
with open(OUT, "w") as out:
    for r in (target_cs, round_cs):
        out.write(json.dumps(r) + "\n")
    for lid in sorted(values, key=lambda l: (l in districts, l != "uga", l)):
        for col in COLUMNS:
            out.write(json.dumps(group(lid, col)) + "\n")
            n += 1
nat = values["uga"]
print(f"{n} Groups ({len(values)} places × {len(COLUMNS)} targets) and 2 CodeSystems → {OUT}")
for col, (code, *_rest) in COLUMNS.items():
    print(f"  {code:26} Uganda {round(nat[col]):>12,}")
